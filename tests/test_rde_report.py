import asyncio
import os
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest

from pptx import Presentation
from pydantic import ValidationError
from app.report import summarize, WorkItem, Report
from app.ppt import render, frames


class StructuredReportTests(unittest.TestCase):
    def cfg(self):
        return SimpleNamespace(mode="mock", max_team=30000)

    def sample(self):
        return [{"user_id": 1, "name": "최찬", "week": "2026-10-19", "body":
            "[금주 업무]\n"
            "사업 | 도로공사 | ENV 연결 확인 진행 중 | 최찬 | 진행 | 2026-10-22\n"
            "내부 | 서버 운영 | 메모리 확인 | 최찬 | 완료 | 2026-10-21\n"
            "[차주 계획]\n사업 | 도로공사 | Studio 기능 테스트 | 최찬 | 미기재 | 다음 주\n"
            "[외근 일정]\n2026-10-27 10:00 | 도로공사 | 연결 시험 | 최찬 | 조율 중 | 출입 확인\n"
            "[이슈]\n도로공사 | DNS 등록 대기 | 담당자 확인 요청"}]

    def test_structured_mock(self):
        report = asyncio.run(summarize(self.sample(), self.cfg()))
        m = report.members[0]
        self.assertEqual(m.work_items[0].status, "진행")
        self.assertEqual(m.work_items[0].start_date, "2026-10-22")
        self.assertEqual(m.work_items[-1].start_date, "")
        self.assertEqual(m.work_items[1].category, "internal")
        self.assertEqual(m.visits[0].status, "조율 중")
        self.assertIn("Studio 기능 테스트", m.next_week)

    def test_missing_fields_not_invented(self):
        rows = [{"user_id": 1, "name": "최찬", "body": "[금주 업무]\n- 환경 구성 진행 중"}]
        report = asyncio.run(summarize(rows, self.cfg()))
        task = report.members[0].work_items[0]
        self.assertEqual(task.status, "미기재")
        self.assertEqual(task.project, "미분류")
        self.assertEqual(task.schedule, "")
        self.assertEqual(task.start_date, "")
        self.assertEqual(report.members[0].issue_items, [])

    def test_date_validation(self):
        values = dict(category="business", project="test", phase="actual", description="task",
                      assignee="", status="", schedule="", start_date="2026-02-30", end_date="")
        with self.assertRaises(ValidationError):
            WorkItem(**values)

    def test_schema(self):
        schema = Report.model_json_schema()
        self.assertEqual(schema["additionalProperties"], False)
        self.assertIn("WorkItem", schema["$defs"])

    def test_real_rde_template_and_pagination(self):
        template = Path(os.getenv("RDE_WEEKLY_TEMPLATE", "templates/weekly-template.pptx"))
        if not template.is_file():
            self.skipTest("RDE_WEEKLY_TEMPLATE에 원본 주간 PPTX를 지정하세요.")
        rows = self.sample()
        # More tasks than the original seven data rows; preserve the last item.
        rows[0]["body"] += "\n[금주 업무]\n" + "\n".join(
            f"사업 | 도로공사 | 추가업무-{i:02d} | 최찬 | 진행 | 미정" for i in range(35))
        report = asyncio.run(summarize(rows, self.cfg()))
        with TemporaryDirectory() as tmp:
            out = Path(tmp)/"weekly.pptx"
            render(template, out, report, rows, "2026-10-19", "mock")
            prs = Presentation(out)
            text = "\n".join(f.text for s in prs.slides for f in frames(s.shapes))
            self.assertGreater(len(prs.slides), 3)
            self.assertIn("추가업무-34", text)
            self.assertIn("ENV 연결 확인 진행 중", text.replace("\n", ""))
            self.assertIn("DNS 등록 대기", text)
            self.assertIn("조율 중", text)
            self.assertNotIn("우리은행", text)
            self.assertNotIn("SANCTION", text)
            self.assertNotIn("7월9일", text)
            self.assertNotIn("작성 예시", text)
            for slide in prs.slides:
                for sh in slide.shapes:
                    if sh.has_table:
                        self.assertLessEqual(sh.top + sh.height, prs.slide_height - 200000)


if __name__ == "__main__":
    unittest.main()
