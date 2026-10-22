import asyncio
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from zoneinfo import ZoneInfo
from pptx import Presentation
from pptx.util import Inches, Pt
from app.db import DB
from app.report import week_start, summarize
from app.ppt import render, inspect_template, replace_paragraph


class CoreTests(unittest.TestCase):
    def test_input_tokens_are_literal(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        p = slide.shapes.add_textbox(0, 0, Inches(8), Inches(4)).text_frame.paragraphs[0]
        for text in ['{{COM', 'PLETED}} + {{ISSUES}}']:
            p.add_run().text = text
        replace_paragraph(p, {'{{COMPLETED}}': '{{ISSUES}} literal', '{{ISSUES}}': 'issue'})
        self.assertEqual(p.text, '{{ISSUES}} literal + issue')

    def test_scope_replace_and_old_update(self):
        with TemporaryDirectory() as tmp:
            db = DB(Path(tmp)/'test.db')
            db.submit(-1, '2026-10-19', 1, '찬', 'old', 10)
            db.submit(-1, '2026-10-19', 1, '찬', 'new', 12)
            db.submit(-1, '2026-10-19', 1, '찬', 'replayed', 11)
            self.assertEqual(db.list(-1, '2026-10-19')[0]['body'], 'new')
            self.assertEqual(db.list(-2, '2026-10-19'), [])
            self.assertEqual(db.list(-1, '2026-10-26'), [])
            self.assertTrue(db.job('job1', -1, '2026-10-19'))
            self.assertFalse(db.job('job1', -1, '2026-10-19'))
            DB(Path(tmp)/'test.db')
            with db.connect() as c:
                self.assertEqual(c.execute('SELECT state FROM jobs').fetchone()[0], 'interrupted')

    def test_week_boundary(self):
        self.assertEqual(week_start('2026-10-25', ZoneInfo('Asia/Seoul')), '2026-10-19')
        self.assertEqual(week_start('2027-01-01'), '2026-12-28')
        with self.assertRaises(ValueError):
            week_start('2026-02-30')

    def test_mock_and_template_roundtrip(self):
        rows = [{'user_id': 1, 'name': '최찬', 'body': '[금주 업무]\n- 연결 확인 진행 중\n[차주 계획]\n- 테스트\n[이슈]\n- DNS 대기'}]
        cfg = SimpleNamespace(mode='mock', max_team=30000)
        report = asyncio.run(summarize(rows, cfg))
        self.assertEqual(report.members[0].completed, ['연결 확인 진행 중'])
        with TemporaryDirectory() as tmp:
            # 테스트용 최소 PPT fixture. 배포할 샘플 PPT가 아님
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4)).text_frame
            p = frame.paragraphs[0]
            for text in ['보고 {{COM', 'PLETED}} 끝']:
                run = p.add_run(); run.text = text; run.font.size = Pt(20)
            frame.add_paragraph().text = '{{NEXT_WEEK}}'
            frame.add_paragraph().text = '{{ISSUES}}'
            template, out = Path(tmp)/'template.pptx', Path(tmp)/'out.pptx'
            prs.save(template)
            self.assertIn('{{COMPLETED}}', inspect_template(template))
            render(template, out, report, rows, '2026-10-19', 'mock')
            actual = Presentation(out).slides[0].shapes[0].text
            self.assertIn('최찬', actual)
            self.assertIn('진행 중', actual)
            self.assertNotIn('{{', actual)
            self.assertIn('끝', actual)
            self.assertEqual(Presentation(out).slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size.pt, 20)

    def test_missing_template_and_limit(self):
        cfg = SimpleNamespace(mode='mock', max_team=1)
        with self.assertRaises(ValueError):
            asyncio.run(summarize([{'body': 'long'}], cfg))
        with self.assertRaises(ValueError):
            render(Path('/not/a/template.pptx'), Path('/tmp/out.pptx'), None, [], '2026-10-19', 'mock')


if __name__ == '__main__':
    unittest.main()
