from datetime import date, timedelta
import re
from pptx import Presentation

TOKENS = {'{{TITLE}}', '{{PERIOD}}', '{{COMPLETED}}', '{{NEXT_WEEK}}', '{{ISSUES}}', '{{MEMBERS}}'}


def frames(shapes):
    for shape in shapes:
        if shape.shape_type == 6:  # GROUP
            yield from frames(shape.shapes)
        if shape.has_text_frame:
            yield shape.text_frame
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def inspect_template(path):
    prs = Presentation(str(path))
    found = set()
    for slide in prs.slides:
        for frame in frames(slide.shapes):
            found.update(re.findall(r'\{\{[A-Z_]+\}\}', frame.text))
    return found


def replace_paragraph(p, values):
    # 토큰이 여러 run에 나뉘어 있어도 교체. 주변 run과 첫 run의 스타일을 유지
    # 원문에 포함된 토큰 모양의 문자열을 다시 템플릿으로 해석하지 않음
    runs = list(p.runs)
    text = ''.join(r.text for r in runs)
    pattern = '|'.join(re.escape(t) for t in values)
    for match in reversed(list(re.finditer(pattern, text))):
        start, end = match.span()
        value = values[match.group()]
        offset = 0
        for r in runs:
            old = r.text
            a, b = offset, offset + len(old)
            if b > start and a < end:
                prefix = old[:max(0, start-a)]
                suffix = old[max(0, end-a):] if end < b else ''
                r.text = prefix + (value if a <= start < b else '') + suffix
            offset = b


def render(template, output, report, rows, week, mode):
    if not template.is_file():
        raise ValueError('templates/weekly-template.pptx 파일을 준비해 주세요.')
    found = inspect_template(template)
    required = {'{{COMPLETED}}', '{{NEXT_WEEK}}', '{{ISSUES}}'}
    if not required.issubset(found):
        raise ValueError('PPT 양식에 {{COMPLETED}}, {{NEXT_WEEK}}, {{ISSUES}}가 필요합니다.')
    if found - TOKENS:
        raise ValueError('지원하지 않는 PPT 토큰이 있습니다.')
    names = {str(r['user_id']): r['name'] for r in rows}
    def section(key):
        result = []
        for member in report.members:
            items = getattr(member, key)
            result.append(names[member.user_id] + '\n' + ('\n'.join('• ' + x for x in items) if items else '미기재'))
        return '\n\n'.join(result)
    start = date.fromisoformat(week)
    values = {
        '{{TITLE}}': '주간업무보고' + (' (원문 테스트)' if mode == 'mock' else ' (검토용 초안)'),
        '{{PERIOD}}': f'{start:%Y.%m.%d} ~ {start+timedelta(days=6):%Y.%m.%d}',
        '{{MEMBERS}}': ', '.join(names.values()),
        '{{COMPLETED}}': section('completed'), '{{NEXT_WEEK}}': section('next_week'), '{{ISSUES}}': section('issues'),
    }
    prs = Presentation(str(template))
    for slide in prs.slides:
        for frame in frames(slide.shapes):
            for paragraph in frame.paragraphs:
                replace_paragraph(paragraph, values)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    # XML 검사는 시각적 넘침을 보장하지 않음. 원문/구조화 결과를 함께 보관
    Presentation(str(output))
    return output
